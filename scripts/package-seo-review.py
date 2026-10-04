"""Generate only the four requested handoff documents; no hosting operations."""
import hashlib
import json
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs/seo/handoff-simple-20261004'
OUTPUT = ROOT / '.astro/reports'
REPORT = json.loads((SOURCE / 'report.json').read_text())
OUTPUT.mkdir(parents=True, exist_ok=True)
workbook = Workbook()
summary = workbook.active
summary.title = 'Sažetak'
summary.append(['DENTVITALIS — SEO MIGRACIJA', 'Pregled za SEO mastera i webmastera'])
summary.append(['Datum', '4. listopada 2026.'])
summary.append(['Status', 'Paket za pregled. Nije instaliran na cPanel.'])
summary.append(['Stari sitemap', '197 URL-ova; original s www.dentvitalis.com, bez izmjena.'])
summary.append(['Novi sitemap', '136 URL-ova: IT 28, HR 27, DE 27, EN 27, SL 27.'])
summary.append(['Hreflang', '812 XML veza; it/hr/de/en/sl + x-default. Provjereni prema HTML-u.'])
summary.append(['Ista adresa', 46])
summary.append(['Postojeća/odobrena 301', 8])
summary.append(['Dodatna 301 u ovom prijedlogu', 8])
summary.append(['Izričito očuvane PHP stranice', 3])
summary.append(['Zahvalne stranice', '5; noindex, izvan novog sitemapa.'])
summary.append(['Ostale stare adrese', '127: bez automatskog preusmjeravanja; razlog i mogući cilj navedeni u usporedbi.'])
summary.append(['Što to znači', '106 je u ranijem inventaru vraćalo 200, a 21 već 404. PHP fallback je očuvan; budući javni prikaz nije ovdje dokazan.'])
summary.append(['.htaccess', '64 naslijeđene 301 pojave uz odobrene ispravke + 16 novih pravila (8 ranije odobrenih, 8 za pregled).'])
summary.append(['Što se ne radi', 'Nema slanja svih starih URL-ova na naslovnicu, lažnih jezičnih parova niti proizvoljnog spajanja medicinskih usluga.'])
summary.append(['Sitemap iznimka', 'Talijanski prijevoz ima samo IT i x-default: ostali prijevodi ne postoje.'])
summary.append(['Prije objave', 'SEO/vlasnik potvrđuju dodatne 301 i sadržajne odluke. Webmaster priprema novu reviziju manifesta i radi HTTP provjeru na hostingu.'])
summary.append(['VAŽNO', 'Uploadani localized-v3 i pripremljeni next root još NE sadrže ove nove SEO datoteke. Ne aktivirati ih kao da sadrže ovu reviziju.'])
summary.append(['Četiri dokumenta', 'sitemap-stari.xml; sitemap-novi.xml; .htaccess; DentVitalis-SEO-usporedba-20261004.xlsx'])

labels = {
    'same-url': 'Ista adresa',
    'existing-or-approved-301': 'Postojeća/odobrena 301',
    'equivalent-301-in-review-candidate': 'Nova 301 u prijedlogu',
    'preserved-PHP-runtime-check-required': 'Očuvana PHP stranica — HTTP provjera',
    'preserved-static-noindex-not-in-new-sitemap': 'Zahvalna stranica — noindex',
    'content-decision-required-no-new-redirect-invented': 'Bez nove 301 — sadržajna odluka',
    'approved-410': 'Odobreno 410',
    'slash-normalization-308': 'Ista stranica — slash 308',
}

def data_sheet(name, headings, rows, widths, table_name):
    sheet = workbook.create_sheet(name)
    sheet.append(headings)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'C2'
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.row_dimensions[1].height = 30
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='Calibri', size=11)
            if isinstance(cell.value, str) and cell.value.startswith('https://') and ' | ' not in cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(name='Calibri', size=11, color='0563C1')
        sheet.row_dimensions[row[0].row].height = 62
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = 'DentVitalis | SEO pregled | &P / &N'
    return sheet

comparison = data_sheet('Stari → novi',
    ['Jezik', 'Status', 'Stari URL', 'Novi URL / aktivno pravilo', 'Mogući cilj — treba odluku', 'Napomena'],
    [[row['language'].upper(), labels[row['status']], row['old'], row['target'],
      '' if row['target'] else ' | '.join(row['proposalTargets']), row['reason']]
     for row in REPORT['comparison']],
    [9, 33, 65, 65, 65, 75], 'StariNovi')
for row in comparison.iter_rows(min_row=2):
    color = 'FFF2CC' if row[1].value.startswith('Bez nove') else 'E2F0D9'
    row[1].fill = PatternFill('solid', fgColor=color)

import xml.etree.ElementTree as ET
ns = {'s': 'http://www.sitemaps.org/schemas/sitemap/0.9', 'h': 'http://www.w3.org/1999/xhtml'}
new_tree = ET.parse(SOURCE / 'sitemap-novi.xml')
new_rows = []
for entry in new_tree.getroot().findall('s:url', ns):
    url = entry.findtext('s:loc', namespaces=ns)
    links = entry.findall('h:link', ns)
    own = next(link.attrib['hreflang'] for link in links if link.attrib['href'] == url and link.attrib['hreflang'] != 'x-default')
    new_rows.append([own.upper(), url, ', '.join(link.attrib['hreflang'] for link in links)])
data_sheet('Novi sitemap', ['Jezik', 'Novi URL', 'Hreflang'], new_rows, [10, 100, 35], 'NoviSitemap')

import re
htaccess = (SOURCE / '.htaccess').read_text()
redirect_rows = []
for source, target in re.findall(r'^Redirect\s+301\s+(\S+)\s+(\S+)\s*$', htaccess, re.M):
    redirect_rows.append(['Staro pravilo (očuvano/ispravljeno)', 301, source, target])
additional = {item['from'] for item in REPORT['reviewRedirects']}
for source, target in re.findall(r'^RewriteRule \^([a-z0-9/-]+)/\?\$ (\S+) \[R=301,L,NE\]$', htaccess, re.M):
    source = '/' + source
    redirect_rows.append(['Novo za pregled' if source in additional else 'Ranije odobreno', 301, source, target])
data_sheet('Redirekcije', ['Vrsta', 'HTTP', 'Stari put', 'Novi put'], redirect_rows, [40, 10, 80, 80], 'Redirekcije')

summary.column_dimensions['A'].width = 38
summary.column_dimensions['B'].width = 115
summary.freeze_panes = 'B3'
summary.sheet_view.showGridLines = False
for row in summary.iter_rows():
    summary.row_dimensions[row[0].row].height = 36
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.font = Font(name='Calibri', size=11, bold=cell.column == 1)
for cell in summary[1]:
    cell.fill = PatternFill('solid', fgColor='17365D')
    cell.font = Font(name='Calibri', color='FFFFFF', bold=True, size=14)
summary.row_dimensions[1].height = 40
for row in [3, 17, 18]:
    for cell in summary[row]:
        cell.fill = PatternFill('solid', fgColor='FFF2CC')

xlsx_name = 'DentVitalis-SEO-usporedba-20261004.xlsx'
xlsx = SOURCE / xlsx_name
workbook.save(xlsx)
checked = load_workbook(xlsx)
assert checked['Stari → novi'].max_row == 198
assert checked['Novi sitemap'].max_row == 137
assert checked['Redirekcije'].max_row == 81
assert len(checked.sheetnames) == 4

zip_path = OUTPUT / 'DentVitalis-SEO-ispravljeno-20261004.zip'
with zipfile.ZipFile(zip_path, 'x', zipfile.ZIP_DEFLATED) as archive:
    for name in ['sitemap-stari.xml', 'sitemap-novi.xml', '.htaccess']:
        archive.write(SOURCE / name, name)
    archive.write(xlsx, xlsx_name)
with zipfile.ZipFile(zip_path) as archive:
    assert archive.testzip() is None
    assert len(archive.namelist()) == 4
    for name in ['sitemap-stari.xml', 'sitemap-novi.xml', '.htaccess']:
        assert archive.read(name) == (SOURCE / name).read_bytes()
    assert archive.read(xlsx_name) == xlsx.read_bytes()
print(json.dumps({'zip': str(zip_path), 'bytes': zip_path.stat().st_size,
                  'sha256': hashlib.sha256(zip_path.read_bytes()).hexdigest(),
                  'documents': 4, 'oldURLs': 197, 'newURLs': 136,
                  'hreflangLinks': 812, 'redirectLines': len(redirect_rows)}, indent=2))
